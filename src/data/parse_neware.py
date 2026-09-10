"""
Parse Neware "intracycle" xlsx exports for VL coin-cell battery tests.

Expected columns (in order):
    Battery type, DataPoint, Cycle Index, Step Index, Step Type,
    Elapsed Time(s), Temperature, Time, Total Time,
    Current(A), Voltage(V), Capacity(Ah), SOC/DOD(%)

Each file is a single sheet named "record" containing one cell's full
cycling history at fixed ~30s sampling. Files may have trailing fully-blank
rows (observed in practice) which must be stripped.

This module only handles reading + cleaning; SOH labeling and feature
extraction live in soh_features.py.
"""
from __future__ import annotations

import re
import warnings
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

EXPECTED_COLUMNS = [
    "Battery type",
    "DataPoint",
    "Cycle Index",
    "Step Index",
    "Step Type",
    "Elapsed Time(s)",
    "Temperature",
    "Time",
    "Total Time",
    "Current(A)",
    "Voltage(V)",
    "Capacity(Ah)",
    "SOC/DOD(%)",
]

# Canonical internal column names (snake_case, unit-free) used everywhere
# downstream of this parser.
COLUMN_RENAME = {
    "Battery type": "battery_type",
    "DataPoint": "datapoint",
    "Cycle Index": "cycle",
    "Step Index": "step_index",
    "Step Type": "step_type",
    "Elapsed Time(s)": "elapsed_s",
    "Temperature": "temperature_c",
    "Time": "step_time_str",
    "Total Time": "total_time_str",
    "Current(A)": "current_a",
    "Voltage(V)": "voltage_v",
    "Capacity(Ah)": "capacity_ah",
    "SOC/DOD(%)": "soc_dod_pct",
}

# Known battery-type tag -> (form_factor, c_rate, temperature_condition)
# Parsed from the filename / battery_type column, e.g. "VL2020-0.2C-40".
_TAG_RE = re.compile(
    r"^(?P<form_factor>VL\d{4})-(?P<c_rate>[\d.]+C)-(?P<temp_tag>RT|\d+)$"
)


@dataclass
class VLCellMeta:
    cell_id: str          # e.g. "VL2020-0.2C-RT"
    form_factor: str       # e.g. "VL2020"
    c_rate: str            # e.g. "0.2C"
    temp_condition: str    # "RT" or e.g. "40" (deg C nominal)

    @classmethod
    def from_tag(cls, tag: str) -> "VLCellMeta":
        m = _TAG_RE.match(tag)
        if not m:
            raise ValueError(f"Battery type tag '{tag}' does not match expected pattern")
        return cls(
            cell_id=tag,
            form_factor=m.group("form_factor"),
            c_rate=m.group("c_rate"),
            temp_condition=m.group("temp_tag"),
        )


def read_neware_xlsx(path: str | Path) -> pd.DataFrame:
    """Read a single Neware intracycle xlsx export into a cleaned DataFrame.

    - Validates the header matches the expected schema.
    - Drops fully-blank trailing rows.
    - Coerces dtypes.
    - Sorts by DataPoint to guarantee chronological order.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_name = "record" if "record" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    header = [str(h).strip() if h is not None else h for h in header]
    if header[: len(EXPECTED_COLUMNS)] != EXPECTED_COLUMNS:
        raise ValueError(
            f"{path.name}: header mismatch.\n  expected={EXPECTED_COLUMNS}\n  got={header}"
        )

    data = [r for r in rows if r[2] is not None]  # drop blank trailer rows
    wb.close()

    df = pd.DataFrame(data, columns=EXPECTED_COLUMNS[: len(data[0])])
    df = df.rename(columns=COLUMN_RENAME)

    # dtypes
    int_cols = ["datapoint", "cycle", "step_index"]
    float_cols = ["elapsed_s", "temperature_c", "current_a", "voltage_v", "capacity_ah", "soc_dod_pct"]
    for c in int_cols:
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    for c in float_cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["step_type"] = df["step_type"].astype("category")
    df["battery_type"] = df["battery_type"].astype(str)

    df = df.sort_values("datapoint").reset_index(drop=True)

    # sanity checks
    # Temperature sensor dropouts are observed in practice (isolated gaps of
    # a handful of consecutive rows where every other column is valid). We
    # repair these with linear interpolation rather than dropping the rows,
    # since dropping would create a hole in the capacity/voltage trace for
    # that step. Any NaN in a column other than temperature is fatal.
    critical_cols = ["datapoint", "cycle", "step_index", "elapsed_s", "current_a", "voltage_v", "capacity_ah", "soc_dod_pct"]
    n_bad_critical = df[critical_cols].isna().any(axis=1).sum()
    if n_bad_critical:
        raise ValueError(f"{path.name}: {n_bad_critical} rows have unexpected NaNs in critical columns")

    n_temp_nan = df["temperature_c"].isna().sum()
    if n_temp_nan:
        max_gap = (
            df["temperature_c"].isna().astype(int).groupby(
                (~df["temperature_c"].isna()).cumsum()
            ).sum().max()
        )
        if max_gap > 60:
            raise ValueError(
                f"{path.name}: temperature gap of {max_gap} consecutive rows is too "
                "large to safely interpolate"
            )
        warnings.warn(
            f"{path.name}: interpolating {n_temp_nan} missing temperature reading(s) "
            f"(longest gap: {max_gap} rows)"
        )
        df["temperature_c"] = df["temperature_c"].interpolate(method="linear", limit_direction="both")

    tags = df["battery_type"].unique()
    if len(tags) != 1:
        raise ValueError(f"{path.name}: expected a single battery_type tag, found {tags}")

    return df


def infer_meta(df: pd.DataFrame) -> VLCellMeta:
    tag = df["battery_type"].iloc[0]
    return VLCellMeta.from_tag(tag)


def validate_cycle_structure(df: pd.DataFrame) -> dict:
    """Return a small diagnostics dict useful for spotting anomalies:
    step-type sequence per cycle, row counts, cycle count, gaps.
    """
    cycles = sorted(df["cycle"].unique())
    step_types = sorted(df["step_type"].dropna().unique().tolist())
    rows_per_cycle = df.groupby("cycle").size()
    # expected repeating pattern after cycle 1: CC DChg -> Rest -> CC Chg -> CV Chg -> Rest
    seq_cycle2 = (
        df[df["cycle"] == cycles[min(1, len(cycles) - 1)]]
        .drop_duplicates(subset="step_index")["step_type"]
        .tolist()
    )
    return {
        "n_cycles": len(cycles),
        "cycle_min": int(cycles[0]),
        "cycle_max": int(cycles[-1]),
        "step_types": step_types,
        "rows_per_cycle_min": int(rows_per_cycle.min()),
        "rows_per_cycle_max": int(rows_per_cycle.max()),
        "sample_step_sequence": seq_cycle2,
        "missing_cycles": sorted(set(range(cycles[0], cycles[-1] + 1)) - set(cycles)),
    }


if __name__ == "__main__":
    import sys
    import json

    for path in sys.argv[1:]:
        df = read_neware_xlsx(path)
        meta = infer_meta(df)
        diag = validate_cycle_structure(df)
        print(f"\n=== {path} ===")
        print("meta:", meta)
        print("diagnostics:", json.dumps(diag, default=str, indent=2))
